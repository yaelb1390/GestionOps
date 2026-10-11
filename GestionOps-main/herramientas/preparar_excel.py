#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Prepara un Excel de GestionOps para importarlo en el panel (Importar Excel).

Lee el archivo tal como lo exporta el sistema, detecta el módulo (Tickets, Órdenes, Averías repetidas o
Códigos razón), ordena y limpia los datos con las mismas reglas que usa el panel, y escribe un Excel nuevo
con SOLO las columnas de ese módulo, listo para subir.

Uso:
    pip install openpyxl
    python preparar_excel.py ordenes.xlsx
    python preparar_excel.py ordenes.xlsx --modulo ordenes -o ordenes_listo.xlsx
    python preparar_excel.py *.xlsx                      # varios archivos a la vez

Opciones:
    --modulo {auto,tickets,ordenes,calidad,razones}   (por defecto: auto)
    -o / --salida ARCHIVO     solo con un archivo de entrada (por defecto: <nombre>_listo.xlsx)
    --hoja NOMBRE             hoja a leer (por defecto: la primera que tenga datos)
    --con-extras              conserva también las demás columnas del archivo, al final

Qué hace:
  · Encuentra la fila de encabezados aunque haya títulos encima.
  · Si dos columnas significan lo mismo (p. ej. "Trabajo" y "Work Name"), gana "Trabajo"; la otra se guarda
    aparte con su propio nombre (con --con-extras) y no se pierde.
  · En Órdenes, "Orden externa" es la clave y "Trabajo" siempre es la descripción.
  · Las fechas (texto "13/10/2026 12:00 p. m.", números de Excel, "2026-10-13 09:30", "15-oct-2026"...) se
    convierten a fechas reales, hora de República Dominicana.
  · Los identificadores se guardan como texto, sin ".0" ni notación científica.
  · Quita filas vacías, sin identificador y repetidas (quedan listadas en la hoja "Descartadas").
  · No incluye estado, inspector ni código aplicado: el panel nunca los pisa.
"""
import argparse
import csv
import datetime as dt
import re
import sys
import unicodedata
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover
    sys.exit("Falta la librería openpyxl. Instálala con:  pip install openpyxl")

# ── Módulos: columnas de salida (clave interna, encabezado que el panel reconoce) ────────────────────────────
MODULOS = {
    "tickets": [("ticket", "Ticket"), ("descripcion_orden", "Descripcion"), ("supervisor", "Supervisor"),
                ("tech", "Nombre Tecnico"), ("sector", "Sector"), ("tecnologia", "Tecnologia"),
                ("priority", "Prioridad"), ("fecha", "Fecha")],
    "ordenes": [("orden_externa_id", "Orden externa"), ("descripcion_orden", "Trabajo"), ("cliente", "Cliente"),
                ("tecnologia", "Tecnologia"), ("sector", "Sector"), ("terminal", "Terminal"),
                ("supervisor", "Supervisor"), ("tech", "Nombre Tecnico"), ("oe_vence", "OE Vencimiento"),
                ("priority", "Prioridad")],
    "calidad": [("ticket", "Caso actual"), ("work_name", "Work name actual"), ("caso_repetido", "Caso repetido"),
                ("tecnico_repetido", "Tecnico caso repetido"), ("respuesta_repetido", "Respuesta caso repetido"),
                ("fecha_repetido", "Fecha cierre repetido"), ("tech", "Nombre Tecnico"),
                ("supervisor", "Supervisor"), ("sector", "Sector")],
    "razones": [("ticket", "Ticket"), ("descripcion_orden", "Descripcion"), ("cliente", "Cliente"),
                ("supervisor", "Supervisor"), ("tech", "Nombre Tecnico"), ("sector", "Sector"),
                ("tecnologia", "Tecnologia"), ("priority", "Prioridad"), ("fecha", "Fecha")],
}
NOMBRES = {"tickets": "Tickets", "ordenes": "Órdenes", "calidad": "Averías repetidas", "razones": "Códigos razón"}
# Columnas reales que existen en la base pero no se muestran: se incluyen solo si traen datos
OPCIONALES = [("orden_servicio", "Orden servicio"), ("tech_id", "ID Tecnico"), ("supervisor_id", "ID Supervisor"),
              ("barrio", "Barrio"), ("ciudad", "Ciudad")]
COLS_CALIDAD = {"work_name", "caso_repetido", "tecnico_repetido", "respuesta_repetido", "fecha_repetido"}
FECHAS = {"fecha", "oe_vence", "fecha_repetido"}
PROTEGIDAS = {"status", "inspector", "inspector_id", "codigo_aplicado", "fecha_inspeccion"}
DESC_WORDS = ("reparacion", "instalacion", "averia", "proactiva", "internet", "voip", "cobre")
MESES = {"ene": 1, "jan": 1, "feb": 2, "mar": 3, "abr": 4, "apr": 4, "may": 5, "jun": 6, "jul": 7, "ago": 8,
         "aug": 8, "sep": 9, "set": 9, "oct": 10, "nov": 11, "dic": 12, "dec": 12}
RD = dt.timezone(dt.timedelta(hours=-4))


# ── Normalización (mismas reglas que el panel y que el antiguo Apps Script) ──────────────────────────────────
def norm(s):
    if s is None:
        return ""
    s = unicodedata.normalize("NFD", str(s).strip().lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def snake(s):
    return re.sub(r"[^a-z0-9_]", "", re.sub(r"\s+", "_", norm(s)))


def normalize_header(h):
    low = norm(h)
    if not low:
        return ""
    def es(*xs):
        return low in xs
    if es("trabajo", "idd", "nro trabajo", "id trabajo", "id_trabajo", "job id", "nro orden", "numero de orden",
          "nro de orden", "nro_orden", "ticket", "id ticket", "nro ticket", "id_ticket"):
        return "ticket"
    if es("work name", "tipo trabajo", "tipo de trabajo", "descripcion trabajo", "tipo servicio", "tipo_servicio"):
        return "descripcion_orden"
    if es("orden servicio", "orden de servicio", "orden_servicio", "os", "id orden"):
        return "orden_servicio"
    if es("orden externa", "order externa", "external order"):
        return "orden_externa_id"
    if es("descripcion", "descripcion orden"):
        return "descripcion_orden"
    if es("cliente", "nombre cliente", "nombre_cliente", "subscriber"):
        return "cliente"
    if es("fecha", "fecha reporte", "fecha_creacion"):
        return "fecha"
    if es("vence", "oe vencimiento", "oe_vencimiento", "oe vence", "oe_vence"):
        return "oe_vence"
    if es("tecnologia", "tipo red", "tipo_red", "red", "servicio", "tipo servicio"):
        return "tecnologia"
    if low == "sector":
        return "sector"
    if low == "barrio":
        return "barrio"
    if low == "ciudad":
        return "ciudad"
    if es("zona", "region", "municipio"):
        return "sector"
    if es("terminal", "id terminal", "fat", "caja"):
        return "terminal"
    if "supervisor" in low and "id" not in low and "tarjeta" not in low:
        return "supervisor"
    if "supervisor" in low and ("id" in low or "tarjeta" in low):
        return "supervisor_id"
    if es("nombre tecnico", "nombre_tecnico", "asignado a", "asignado_a", "tech_name", "tech"):
        return "tech"
    if es("tecnico", "cedula", "cedula tecnico", "tarjeta") or "id tecnico" in low or "tarjeta tecnico" in low \
            or "tech_id" in low:
        return "tech_id"
    if es("inspector", "inspector_nombre", "nombre inspector", "inspector asignado", "inspector_asignado",
          "nombre_inspector"):
        return "inspector"
    if es("inspector id", "inspector_id", "tarjeta inspector", "id inspector", "id_inspector"):
        return "inspector_id"
    if es("caso actual", "caso_actual", "ticket_actual"):
        return "ticket"
    if es("tecnico actual", "tecnico_actual"):
        return "tech"
    if es("supervisor actual", "supervisor_actual"):
        return "supervisor"
    if es("work name actual", "work_name_actual", "work name"):
        return "work_name"
    if es("caso repetido", "caso_repetido", "ticket_repetido"):
        return "caso_repetido"
    if es("tecnico caso repetido", "tecnico_caso_repetido", "tecnico repetido"):
        return "tecnico_repetido"
    if es("respuesta caso repetido", "respuesta_caso_repetido", "respuesta_repetido"):
        return "respuesta_repetido"
    if es("fecha cierre repetido", "fecha_cierre_repetido", "fecha_repetido"):
        return "fecha_repetido"
    if es("tecnologia actual", "tecnologia_actual"):
        return "tecnologia"
    if es("sector actual", "sector_actual"):
        return "sector"
    if es("estado", "status", "estado_orden", "estado inspeccion", "estado_inspeccion"):
        return "status"
    if es("prioridad", "priority", "tipo de prioridad", "clasificacion prioridad"):
        return "priority"
    if es("codigo aplicado", "codigo_aplicado", "manual_code", "codigo_razon"):
        return "codigo_aplicado"
    if es("fecha inspeccion", "fecha_inspeccion", "fecha registro", "fecha_registro"):
        return "fecha_inspeccion"
    return snake(h)


# ── Fechas ───────────────────────────────────────────────────────────────────────────────────────────────────
_FECHA_RE = re.compile(
    r"^(\d{1,4})[/\-. ]+([a-z]{3,}|\d{1,2})[/\-. ]+(\d{1,4})"
    r"(?:[ t,]+(\d{1,2}):(\d{2})(?::(\d{2}))?(?:\.\d+)?\s*(?:([ap])\.?\s?m\.?)?)?$")


def _serial(n):
    if not 20000 < n < 80000:
        return None
    return (dt.datetime(1899, 12, 30) + dt.timedelta(seconds=round(n * 86400))).replace(microsecond=0)


def parse_date(v):
    """Devuelve datetime (hora de reloj de RD, sin zona) o None."""
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.replace(microsecond=0, tzinfo=None)
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _serial(float(v))
    s = str(v).strip()
    if re.fullmatch(r"\d+(\.\d+)?", s):
        return _serial(float(s))
    m = _FECHA_RE.match(norm(s))
    if m:
        a, b, c, h, mi, se, ap = m.groups()
        mes = int(b) if b.isdigit() else MESES.get(b[:3])
        if not mes:
            return None
        if len(a) == 4:
            y, d = int(a), int(c)
        else:
            y, d = (2000 + int(c) if int(c) < 100 else int(c)), int(a)
        hh = int(h or 0)
        if ap:
            hh = hh % 12 + (12 if ap == "p" else 0)
        try:
            return dt.datetime(y, mes, d, hh, int(mi or 0), int(se or 0))
        except ValueError:
            return None
    if re.search(r"t.*(z|[+-]\d{2}:?\d{2})$", s, re.I):  # ISO con zona explícita
        try:
            return dt.datetime.fromisoformat(s.replace("Z", "+00:00").replace("z", "+00:00")) \
                .astimezone(RD).replace(tzinfo=None, microsecond=0)
        except ValueError:
            return None
    return None


# ── Lectura ──────────────────────────────────────────────────────────────────────────────────────────────────
def texto(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.isoformat(sep=" ")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def limpiar_id(v):
    return re.sub(r"\.0$", "", v.strip())


def leer_hojas(ruta, hoja=None):
    """Devuelve (nombre_hoja, filas) con la primera hoja que tenga datos."""
    if ruta.suffix.lower() == ".csv":
        with open(ruta, newline="", encoding="utf-8-sig") as f:
            return "CSV", [list(r) for r in csv.reader(f)]
    if ruta.suffix.lower() == ".xls":
        sys.exit(f"{ruta.name}: el formato .xls antiguo no se puede leer. Ábrelo en Excel y guárdalo como .xlsx")
    wb = load_workbook(ruta, data_only=True)
    nombres = [hoja] if hoja else wb.sheetnames
    for n in nombres:
        if n not in wb.sheetnames:
            sys.exit(f"No existe la hoja '{n}'. Hojas disponibles: {', '.join(wb.sheetnames)}")
        filas = [list(r) for r in wb[n].iter_rows(values_only=True)]
        filas = [r for r in filas if any(texto(c) for c in r)]
        if len(filas) > 1 or hoja:
            return n, filas
    return wb.sheetnames[0], []


def fila_encabezados(filas):
    claves = {"trabajo", "ticket", "orden servicio", "id", "caso", "caso actual", "orden externa"}
    for i, r in enumerate(filas[:10]):
        if any(norm(c) in claves for c in r):
            return i
    return 0


def detectar(raw, nh):
    tiene = nh.__contains__
    base = sum(tiene(k) for k in ("ticket", "supervisor", "tech", "priority", "sector", "tecnologia"))
    razon = any("razon" in norm(h) for h in raw)
    sc = {
        "calidad": 3 * sum(tiene(k) for k in ("caso_repetido", "tecnico_repetido", "respuesta_repetido",
                                               "fecha_repetido")) + (1 if tiene("work_name") else 0),
        "ordenes": 6 * tiene("orden_externa_id") + 4 * tiene("orden_servicio") + 3 * tiene("oe_vence")
                   + tiene("cliente") + tiene("terminal"),
        "tickets": base + (1 if tiene("ticket") else 0),
        "razones": base + (1 if tiene("ticket") else 0) + (4 if razon else 0),
    }
    orden = ["calidad", "ordenes", "razones", "tickets"]
    rank = sorted(orden, key=lambda m: -sc[m])
    top, second = rank[0], rank[1]
    gap = sc[top] - sc[second]
    conf = "alta" if sc[top] >= 6 and gap >= 3 else "media" if sc[top] >= 3 and gap >= 1 else "baja"
    if sc[top] == sc[second] and {top, second} <= {"tickets", "razones"}:
        top = "tickets"
    return top, conf, sc


def clave_de(modulo, r):
    c = lambda k: limpiar_id(str(r.get(k, "") or ""))
    if modulo == "ordenes":
        return c("orden_externa_id") or c("orden_servicio") or c("ticket") or None
    if modulo == "calidad":
        return f"{c('ticket')}|{c('caso_repetido')}" if c("ticket") else None
    return c("ticket") or None


# ── Proceso ──────────────────────────────────────────────────────────────────────────────────────────────────
def procesar(ruta, modulo_pedido, hoja, con_extras):
    nombre_hoja, filas = leer_hojas(ruta, hoja)
    if len(filas) < 2:
        sys.exit(f"{ruta.name}: no hay filas con datos.")
    hi = fila_encabezados(filas)
    raw = [texto(h) for h in filas[hi]]
    nh = [normalize_header(h) for h in raw]
    datos = [(hi + 2 + i, r) for i, r in enumerate(filas[hi + 1:])]  # (nº de fila en Excel, celdas)

    modulo, conf, scores = (modulo_pedido, "elegido", None) if modulo_pedido != "auto" else detectar(raw, nh)
    conocidas = {k for k, _ in MODULOS[modulo]} | {k for k, _ in OPCIONALES} | (COLS_CALIDAD if modulo == "calidad" else set())
    conocidas |= {"remarks", "root_cause"}
    tiene_id_orden = "orden_externa_id" in nh or "orden_servicio" in nh
    # "Trabajo" y "Tipo de prioridad" primero: su valor tiene prioridad sobre otras columnas equivalentes
    orden_cols = sorted(range(len(nh)), key=lambda j: 0 if norm(raw[j]) in ("trabajo", "tipo de prioridad") else 1)

    ok, descartadas, vistas = [], [], {}
    extras_nombres, ignoradas, fechas_malas = [], set(), []
    for nfila, celdas in datos:
        o = {}

        fuente = {}

        def aparte(nombre, val, key):
            base = snake(nombre) or key
            alt, n = base, 2
            while alt == key or alt in o:
                alt = f"{base}_{n}"
                n += 1
            o[alt] = val

        def poner(key, val, nombre):
            # Gana la primera columna con valor; en fechas, la primera que SE PUEDA LEER
            # (una columna con "Mañana" no pisa a una fecha). La otra se guarda aparte, no se pierde.
            if val is None or val == "":
                return
            actual = o.get(key)
            if actual is None or actual == "":
                o[key], fuente[key] = val, nombre
            elif actual != val:
                if key in FECHAS and parse_date(actual) is None and parse_date(val) is not None:
                    aparte(fuente[key], actual, key)
                    o[key], fuente[key] = val, nombre
                else:
                    aparte(nombre, val, key)

        for j in orden_cols:
            key = nh[j]
            if not key:
                continue
            cel = celdas[j] if j < len(celdas) else None
            val = texto(cel)
            rawh = norm(raw[j])
            if key in FECHAS and cel is not None and not isinstance(cel, str):
                val = cel  # conserva el tipo (fecha / número) para convertirlo bien
            if rawh in ("ticket", "nro_ticket", "id_ticket", "id"):
                if texto(val):
                    o["ticket"] = texto(val)
            elif rawh == "trabajo":
                t = texto(val)
                if tiene_id_orden:
                    poner("descripcion_orden", t, raw[j])
                elif not any(w in t.lower() for w in DESC_WORDS):
                    if not o.get("ticket"):
                        o["ticket"] = t
                else:
                    poner("descripcion_orden", t, raw[j])
            elif key == "ticket":
                if not o.get("ticket"):
                    o["ticket"] = texto(val)
            else:
                poner(key, val, raw[j])

        rec, extra = {}, {}
        for k, v in o.items():
            if v is None or v == "":
                continue
            if k in PROTEGIDAS:
                ignoradas.add(k)
            elif k in conocidas:
                if k in FECHAS:
                    d = parse_date(v)
                    if d is None:
                        fechas_malas.append(f'{k}: "{texto(v)}" (fila {nfila})')
                        rec[k] = texto(v)
                    else:
                        rec[k] = d
                elif k.endswith("_id") or k in ("ticket", "orden_servicio"):
                    rec[k] = limpiar_id(texto(v))
                else:
                    rec[k] = texto(v)
            else:
                extra[k] = texto(v)
                if k not in extras_nombres:
                    extras_nombres.append(k)
        clave = clave_de(modulo, rec)
        if not clave:
            descartadas.append((nfila, "Sin identificador", celdas))
        elif clave in vistas:
            descartadas.append((nfila, f"Repetida (igual a la fila {vistas[clave]})", celdas))
        else:
            vistas[clave] = nfila
            ok.append((nfila, rec, extra))

    return dict(ruta=ruta, hoja=nombre_hoja, modulo=modulo, conf=conf, scores=scores, ok=ok, descartadas=descartadas,
                total=len(datos), extras=extras_nombres, ignoradas=sorted(ignoradas), fechas_malas=fechas_malas,
                encabezados=raw, con_extras=con_extras)


# ── Escritura ────────────────────────────────────────────────────────────────────────────────────────────────
ENCAB = PatternFill("solid", fgColor="DA291C")


def escribir(res, salida):
    modulo = res["modulo"]
    cols = list(MODULOS[modulo])
    for k, label in OPCIONALES:  # opcionales solo si hay datos
        if any(rec.get(k) for _, rec, _ in res["ok"]) and k not in {c for c, _ in cols}:
            cols.append((k, label))
    # En Códigos razón la columna "razón" es lo que permite reconocer el módulo: se conserva siempre
    extras = [e for e in res["extras"] if res["con_extras"] or (modulo == "razones" and "razon" in e)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append([label for _, label in cols] + extras)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = ENCAB
        c.alignment = Alignment(vertical="center")
    for _, rec, extra in res["ok"]:
        ws.append([rec.get(k) for k, _ in cols] + [extra.get(k) for k in extras])
    for j, (k, _) in enumerate(cols, start=1):
        for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2):
            for c in cell:
                if k in FECHAS and isinstance(c.value, dt.datetime):
                    c.number_format = "dd/mm/yyyy hh:mm"
                elif isinstance(c.value, str):
                    c.number_format = "@"  # texto: evita notación científica en ids largos
    for j in range(1, ws.max_column + 1):
        ancho = max(len(str(c.value)) if c.value is not None else 0 for c in ws[get_column_letter(j)][:200])
        ws.column_dimensions[get_column_letter(j)].width = min(max(ancho + 3, 12), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    r = wb.create_sheet("Resumen")
    filas = [
        ("Archivo original", res["ruta"].name), ("Hoja leída", res["hoja"]),
        ("Módulo", f'{NOMBRES[modulo]}  (detección: {res["conf"]})'),
        ("Filas leídas", res["total"]), ("Filas listas para importar", len(res["ok"])),
        ("Filas descartadas", len(res["descartadas"])),
        ("Fechas que no se pudieron leer", len(res["fechas_malas"])),
        ("Columnas ignoradas (no se pisan en el panel)", ", ".join(res["ignoradas"]) or "—"),
        ("Otras columnas del archivo", ", ".join(res["extras"]) or "—"),
        ("¿Se conservaron en la hoja Datos?", "Sí, al final" if res["con_extras"] else "No (usa --con-extras)" + ("; se mantiene la columna de razón" if res["modulo"] == "razones" else "")),
    ]
    for a, b in filas:
        r.append([a, b])
    for c in r["A"]:
        c.font = Font(bold=True)
    r.column_dimensions["A"].width = 46
    r.column_dimensions["B"].width = 90
    if res["fechas_malas"]:
        r.append([])
        r.append(["Ejemplos de fechas no leídas", ""])
        for e in res["fechas_malas"][:10]:
            r.append(["", e])

    if res["descartadas"]:
        d = wb.create_sheet("Descartadas")
        d.append(["Fila en el Excel", "Motivo"] + [h or f"Col {i + 1}" for i, h in enumerate(res["encabezados"])])
        for c in d[1]:
            c.font = Font(bold=True)
        for nfila, motivo, celdas in res["descartadas"]:
            d.append([nfila, motivo] + [texto(c) for c in celdas])
        d.column_dimensions["B"].width = 38
    wb.save(salida)


def main():
    ap = argparse.ArgumentParser(description="Prepara un Excel de GestionOps para importarlo en el panel.")
    ap.add_argument("archivos", nargs="+", type=Path)
    ap.add_argument("--modulo", choices=["auto", *MODULOS], default="auto")
    ap.add_argument("-o", "--salida", type=Path)
    ap.add_argument("--hoja")
    ap.add_argument("--con-extras", action="store_true")
    a = ap.parse_args()
    if a.salida and len(a.archivos) > 1:
        ap.error("-o/--salida solo se puede usar con un archivo de entrada")
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass
    for ruta in a.archivos:
        if not ruta.exists():
            print(f"✗ {ruta}: no existe")
            continue
        res = procesar(ruta, a.modulo, a.hoja, a.con_extras)
        salida = a.salida or ruta.with_name(ruta.stem + "_listo.xlsx")
        escribir(res, salida)
        print(f"\n✓ {ruta.name}  →  {salida.name}")
        print(f"  Módulo: {NOMBRES[res['modulo']]} (detección {res['conf']})  ·  hoja «{res['hoja']}»")
        print(f"  Filas: {res['total']} leídas · {len(res['ok'])} listas · {len(res['descartadas'])} descartadas")
        if res["conf"] in ("baja", "media"):
            print("  ⚠ Detección poco segura: revisa el módulo o fuérzalo con --modulo")
        if res["ignoradas"]:
            print(f"  Ignoradas (no se pisan): {', '.join(res['ignoradas'])}")
        if res["extras"]:
            print(f"  Otras columnas: {', '.join(res['extras'])}" + ("" if a.con_extras else "  (no incluidas)"))
        if res["fechas_malas"]:
            print(f"  ⚠ {len(res['fechas_malas'])} fecha(s) sin leer, p. ej. {res['fechas_malas'][0]}")


if __name__ == "__main__":
    main()
